import openpyxl
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Load the Excel file
workbook = openpyxl.load_workbook('C:/Users/Alexander/Downloads/hivdr/CTLC_reporting_tool.xlsx')
#workbook = openpyxl.load_workbook('C:/Users/Alexander/Downloads/hivdr/report_data.xlsx')

sheet = workbook.active

def get_resistance_level(value):
    try:
        value = float(value)
        levels = [(60, "High Level Resistance"), (30, "Intermediate Resistance"), (15, "Low Level Resistance"),
                  (10, "Potential Low Level Resistance"), (0, "Susceptible")]
        for threshold, label in levels:
            if value >= threshold:
                return f" ({label})"
    except ValueError:
        return ""
    return ""

def apply_font_formatting(run):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)

for row in sheet.iter_rows(min_row=2, values_only=True):
    doc = Document()

    # Apply styles to 'Normal'
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    paragraph_format = style.paragraph_format
    paragraph_format.line_spacing = 1.15

    # Header
    header = doc.sections[0].header.paragraphs[0]
    header.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    header_run = header.add_run("Ministry of Health\n")
    header_run.bold = True
    apply_font_formatting(header_run)
    header.add_run().add_picture('C:/Users/Alexander/Downloads/hivdr/images.jpg', width=Inches(0.5))
    header_run = header.add_run("\nHIV Drug Resistance Form\n")
    header_run.bold = True
    apply_font_formatting(header_run)

    # LABORATORY INFORMATION
    p = doc.add_paragraph()
    run = p.add_run("LABORATORY INFORMATION:\n")
    run.bold = True
    run.underline = True
    apply_font_formatting(run)
    run = p.add_run("Name of the testing lab: National Public Health Laboratory\nTel:[lab telephone]  "
                    "Email:[lab email]\nContact Name: [lab contact name]\nPhone Number: [lab telephone]\n"
                    "Email: [lab email]")
    apply_font_formatting(run)

    # PATIENT AND LABORATORY INFORMATION
    p = doc.add_paragraph()
    run = p.add_run("PATIENT AND LABORATORY INFORMATION:\n")
    run.bold = True
    run.underline = True
    apply_font_formatting(run)
    run = p.add_run(f"Unique CTC ID Number: {row[0]}\nCollection Date and time: {row[1]}\nSex: {row[2]}  DOB: {row[3]}  Age: {row[4]}\n"
                    f"Received Date and Time: {row[5]}\nReport date and time: {row[6]}\nReceived From: Facility: {row[7]}   District: {row[8]}  "
                    f"Region: {row[9]}\nRequested By: {row[10]}\nReceived by: {row[11]}   Tested Date: {row[12]}")
    apply_font_formatting(run)

    # ART HISTORY
    p = doc.add_paragraph()
    run = p.add_run("ART HISTORY:\n")
    run.bold = True
    run.underline = True
    apply_font_formatting(run)
    run = p.add_run(f"Analyzer: 3500Xl Genetic Analyzer\nUnique instrument ID: NPHL/INFL/SEQ/73\nSpecimen type: Whole blood√  DBS□  Plasma□\n"
                    f"Specimen quality: suitable√  not suitable □")
    apply_font_formatting(run)
    
    doc.add_page_break()
    heading = doc.add_heading('HIVDR Laboratory Report', level=1)
    heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    heading_run = heading.runs[0]
    heading_run.underline = True
    apply_font_formatting(heading_run)

    pre_gart_paragraph = doc.add_paragraph(f"PreGART viral load: {row[13]}")
    pre_gart_paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    apply_font_formatting(pre_gart_paragraph.runs[0])

    doc.add_paragraph("Test Ordered: HIV DRUG RESISTANCE\nSequence includes PR: 10-99\nSequence includes RT: 1-251\nSequence includes IN: 1-263")
    doc.add_paragraph(f"HIV-1 Subtype ……………………………………………………….…. {row[13]}\nPI Major Resistance Associated Mutations………………. {row[14]}\n"
                      f"PI Accessory Resistance Associated Mutations………… {row[15]}")

    inhibitors = [
        ("PROTEASE INHIBITORS", [
            ("atazanavir/r (ATV/r)", row[16]), ("darunavir/r (DRV/r)", row[17]), ("fosamprenavir/r (FPV/r)", row[18]),
            ("indinavir/r (IDV/r)", row[19]), ("lopinavir/r (LPV/r)", row[20]), ("nelfinavir (NFV)", row[21]),
            ("saquinavir/r (SQV/r)", row[22]), ("tipranavir/r (TPV/r)", row[23])
        ]),
        ("NUCLEOSIDE REVERSE TRANSCRIPTASE INHIBITORS", [
            ("abacavir (ABC)", row[25]), ("zidovudine (AZT)", row[26]), ("stavudine (D4T)", row[27]),
            ("didanosine (DDI)", row[28]), ("emtricitabine (FTC)", row[29]), ("lamivudine (3TC)", row[30]),
            ("tenofovir (TDF)", row[31])
        ]),
        ("NON-NUCLEOSIDE REVERSE TRANSCRIPTASE INHIBITORS", [
            ("doravirine (DOR)", row[33]), ("efavirenz (EFV)", row[34]), ("etravirine (ETR)", row[35]),
            ("nevirapine (NVP)", row[36]), ("rilpivirine (RPV)", row[37])
        ]),
        ("INTEGRASE STRAND TRANSFER INHIBITORS", [
            ("Bictegravir (BIC)", row[40]), ("Cabotegravir (CAB)", row[41]), ("Dolutegravir (DTG)", row[42]),
            ("Elvitegravir (EVG)", row[43]), ("Raltegravir (RAL)", row[44])
        ])
    ]

    for title, drugs in inhibitors:
        p = doc.add_paragraph()
        run = p.add_run(f"{title}:\n")
        run.bold = True
        run.underline = True
        apply_font_formatting(run)
        for drug, value in drugs:
            run = p.add_run(f"{drug}…………………… {value}{get_resistance_level(value)}\n")
            apply_font_formatting(run)

    p = doc.add_paragraph(f"NRTI Resistance Associated Mutations …………………… {row[24]}")
    apply_font_formatting(p.runs[0])
    p = doc.add_paragraph(f"NNRTI Resistance Associated Mutations…………………. {row[32]}")
    apply_font_formatting(p.runs[0])
    p = doc.add_paragraph(f"INSTI Major Resistance Associated Mutations: {row[38]}\n"
                          f"INSTI Accessory Resistance Associated Mutations: {row[39]}")
    apply_font_formatting(p.runs[0])

    comments_section = doc.add_table(rows=1, cols=1)
    comments_section.style = 'Table Grid'
    comments_cell = comments_section.cell(0, 0)
    p = comments_cell.paragraphs[0]
    run = p.add_run("COMMENTS:\n")
    run.bold = True
    run.underline = True
    apply_font_formatting(run)
    run = p.add_run(f"Tested by: {row[46]}\nAuthorized by: {row[47]}        Authorized date and time: {row[48]}\n"
                    f"Date report submitted to CTLC: {row[50]}\nName of submitted to CTLC: {row[49]}\n\n")
    apply_font_formatting(run)

    end_report_paragraph = doc.add_paragraph()
    end_report_paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    end_report_run = end_report_paragraph.add_run("......END OF LABORATORY REPORT.........")
    end_report_run.bold = True
    apply_font_formatting(end_report_run)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    footer_run = footer.add_run("Effective date: 25/07/2022 Revision No.0    ")
    footer_run.font.name = 'Times New Roman'
    footer_run.font.size = Pt(12)
    footer.add_run("Page ").alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    page_field = OxmlElement('w:fldSimple')
    page_field.set(qn('w:instr'), 'PAGE')
    footer._element.append(page_field)
    footer_run = footer.add_run(" of ")
    footer_run.font.name = 'Times New Roman'
    footer_run.font.size = Pt(12)
    total_page_field = OxmlElement('w:fldSimple')
    total_page_field.set(qn('w:instr'), 'NUMPAGES')
    footer._element.append(total_page_field)

    doc.save(f'C:/Users/Alexander/Downloads/hivdr/lab_report_{row[0]}.docx')
